"""Idempotent seed: store types + master catalog from docs/DeliveryGram_Master_Catalog.xlsx.

Run: python -m app.seed [path/to/catalog.xlsx]
Commission/tax rates below are PLACEHOLDERS; set real values in the admin panel before launch.
"""
import sys
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import config, models
from app.db import SessionLocal

PHARMACY_NOTICE = (
    "Medicines marked Rx need a valid doctor's prescription, which a licensed pharmacist will verify. "
    "Do not self-medicate. Restricted medicines are not sold online."
)

# code, workbook name, hindi, needs_rx, time_slots, catalog_mode, commission_bps, tax_bps, notice
STORE_TYPES = [
    ("kirana", "Kirana & Grocery", "किराना और राशन", False, False, "structured", 500, 0, None),
    ("medical", "Medical & Pharmacy", "मेडिकल और दवाई", True, False, "structured", 500, 0, PHARMACY_NOTICE),
    ("dairy", "Dairy & Poultry", "डेयरी और अंडे", False, True, "structured", 300, 0, None),
    ("sweets", "Sweet Mart", "मिठाई की दुकान", False, True, "structured", 700, 0, None),
    ("cloth", "Cloth Store", "कपड़े की दुकान", False, False, "photo_or_text", 1000, 0, None),
    ("footwear", "Footwear", "जूते-चप्पल", False, False, "photo_or_text", 1000, 0, None),
    ("general", "General Store", "जनरल स्टोर", False, False, "structured", 700, 0, None),
    ("restaurant", "Restaurant & Tiffin", "होटल और टिफ़िन", False, True, "fixed_menu", 1500, 0, None),
]
DEFAULT_XLSX = Path(__file__).resolve().parents[2] / "docs" / "DeliveryGram_Master_Catalog.xlsx"


def seed_store_types(db: Session) -> dict[str, models.StoreType]:
    out = {}
    for code, name, hi, rx, slots, mode, comm, tax, notice in STORE_TYPES:
        st = db.scalar(select(models.StoreType).where(models.StoreType.code == code))
        if not st:
            st = models.StoreType(code=code, commission_bps=comm, tax_bps=tax)
            db.add(st)
        st.name, st.name_hi, st.needs_prescription, st.has_time_slots = name, hi, rx, slots
        st.catalog_mode, st.safety_notice = mode, notice
        out[name] = st
    db.flush()
    return out


def import_catalog(db: Session, xlsx: Path = DEFAULT_XLSX) -> int:
    types = seed_store_types(db)
    ws = load_workbook(xlsx, read_only=True, data_only=True)["Master Catalog"]
    rows = ws.iter_rows(values_only=True)
    header = {h: i for i, h in enumerate(next(rows))}
    n = 0
    for r in rows:
        code = r[header["Item ID"]]
        if not code:
            continue
        st = types[r[header["Store Category"]]]
        ci = db.scalar(select(models.CatalogItem).where(models.CatalogItem.code == code))
        if not ci:
            ci = models.CatalogItem(code=code)
            db.add(ci)
        img = r[header["Image File"]] if "Image File" in header else None
        ci.store_type_id = st.id
        ci.subcategory = r[header["Subcategory"]]
        ci.name = r[header["Item (English)"]]
        ci.name_hi = r[header["Item (Hindi)"]]
        ci.unit_type = r[header["Unit Type"]]
        ci.pack_sizes = r[header["Typical Pack Sizes"]]
        ci.rural_demand = r[header["Rural Demand"]]
        ci.time_sensitive = r[header["Time-Sensitive"]] == "Yes"
        ci.requires_prescription = r[header["Prescription Required"]] == "Yes"
        ci.image_path = f"/images/{Path(img).name}" if img else None
        n += 1
    db.commit()
    return n


def ensure_admin(db: Session) -> None:
    """Create the first admin from ADMIN_EMAIL; they log in with an emailed OTP. Role is never self-assignable."""
    if not config.ADMIN_EMAIL:
        return
    user = db.scalar(select(models.User).where(models.User.email == config.ADMIN_EMAIL))
    if not user:
        db.add(models.User(email=config.ADMIN_EMAIL, role="admin", name="Admin"))
    elif user.role != "admin":
        user.role = "admin"
    db.commit()


if __name__ == "__main__":
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    with SessionLocal() as session:
        print(f"Imported {import_catalog(session, path)} catalog items")
        ensure_admin(session)
